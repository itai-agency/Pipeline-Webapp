from openpyxl import load_workbook
from pathlib import Path

xlsx = Path('/home/ubuntu/pipeline-clientes-sinahi/sheets_work/current_sheet.xlsx')
wb = load_workbook(xlsx, data_only=False)
for ws in wb.worksheets:
    print('\nSHEET', ws.title, 'max', ws.max_row, ws.max_column)
    headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
    print('HEADERS:')
    for c,h in enumerate(headers, start=1):
        if h is not None:
            print(c, ws.cell(1,c).coordinate, repr(h))
    hits=[]
    for row in ws.iter_rows():
        for cell in row:
            v=cell.value
            if isinstance(v,str) and any(x in v for x in ['Rojo','rojo','Amarillo','amarillo','Verde','verde']):
                hits.append((cell.coordinate, v))
    print('HITS', len(hits))
    for item in hits[:50]:
        print(item)
