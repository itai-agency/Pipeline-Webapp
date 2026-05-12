from openpyxl import load_workbook
from pathlib import Path

base = Path('/home/ubuntu/pipeline-clientes-sinahi/sheets_work')
# Buscar el xlsx descargado/temporal usado en la inspección previa.
candidates = list(base.glob('*.xlsx')) + list(Path('/home/ubuntu/Downloads').glob('*.xlsx'))
if not candidates:
    raise SystemExit('No se encontró ningún XLSX local para generar TSV.')
source = max(candidates, key=lambda p: p.stat().st_mtime)
wb = load_workbook(source, data_only=True)
mapa = {'Rojo': '🔴', 'Amarillo': '🟡', 'Verde': '🟢'}
configs = {
    'RESPALDO_DIARIO': {'range_cols': (13, 22), 'max_row': 1016, 'sem_cols': {13,15,17,19,21,22}, 'out': 'respaldo_m_v.tsv'},
    'SEMANAL': {'range_cols': (11, 20), 'max_row': 999, 'sem_cols': {11,13,15,17,19,20}, 'out': 'semanal_k_t.tsv'},
    'MENSUAL': {'range_cols': (12, 21), 'max_row': 983, 'sem_cols': {12,14,16,18,20,21}, 'out': 'mensual_l_u.tsv'},
}
for sheet, cfg in configs.items():
    ws = wb[sheet]
    start_col, end_col = cfg['range_cols']
    # Usar hasta la última fila con datos dentro del bloque para no pegar filas vacías innecesarias.
    last = 2
    for r in range(3, cfg['max_row'] + 1):
        if any(ws.cell(r, c).value not in (None, '') for c in range(start_col, end_col + 1)):
            last = r
    lines = []
    for r in range(3, last + 1):
        row = []
        for c in range(start_col, end_col + 1):
            val = ws.cell(r, c).value
            if c in cfg['sem_cols']:
                val = mapa.get(val, val)
            if val is None:
                txt = ''
            elif isinstance(val, float):
                txt = str(val)
            else:
                txt = str(val)
            row.append(txt)
        lines.append('\t'.join(row))
    out_path = base / cfg['out']
    out_path.write_text('\n'.join(lines), encoding='utf-8')
    print(sheet, source, out_path, 'rows', len(lines), 'cols', end_col - start_col + 1)
