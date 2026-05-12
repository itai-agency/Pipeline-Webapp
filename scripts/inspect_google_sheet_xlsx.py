from openpyxl import load_workbook
from pathlib import Path
import json

path = Path('/home/ubuntu/pipeline-clientes-sinahi/sheets_work/data_clientes_2026_copia.xlsx')
wb = load_workbook(path, data_only=False)
summary = {}
for ws in wb.worksheets:
    rows = []
    max_row = min(ws.max_row, 15)
    max_col = min(ws.max_column, 20)
    for r in range(1, max_row + 1):
        row = []
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            row.append(v)
        rows.append(row)
    summary[ws.title] = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'sample': rows,
    }

out = Path('/home/ubuntu/pipeline-clientes-sinahi/sheets_work/sheet_structure_summary.json')
out.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(out)
print(json.dumps({k: {'max_row': v['max_row'], 'max_col': v['max_col']} for k, v in summary.items()}, ensure_ascii=False, indent=2))
