from openpyxl import load_workbook
from pathlib import Path
import json

xlsx = Path('/home/ubuntu/pipeline-clientes-sinahi/sheets_work/validation/google_sheet_editado.xlsx')
formulas_wb = load_workbook(xlsx, data_only=False)
values_wb = load_workbook(xlsx, data_only=True)
summary = {}
for sheet_name in formulas_wb.sheetnames:
    ws_f = formulas_wb[sheet_name]
    ws_v = values_wb[sheet_name]
    max_row = min(ws_f.max_row, 15)
    max_col = min(ws_f.max_column, 18)
    rows = []
    for r in range(1, max_row + 1):
        row = []
        for c in range(1, max_col + 1):
            val = ws_v.cell(r, c).value
            formula = ws_f.cell(r, c).value
            if isinstance(formula, str) and formula.startswith('='):
                row.append({'value': val, 'formula': formula[:250]})
            else:
                row.append(val)
        rows.append(row)
    summary[sheet_name] = {
        'dimensions': {'rows': ws_f.max_row, 'cols': ws_f.max_column},
        'sample': rows,
    }

out = Path('/home/ubuntu/pipeline-clientes-sinahi/sheets_work/validation/edit_validation_summary.json')
out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(out)
